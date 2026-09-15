import openpyxl
from openpyxl import load_workbook
from datetime import datetime
import os
from app.config import EXCEL_PATH

class ExcelExporter:
    def __init__(self, file_path: str = EXCEL_PATH):
        self.file_path = file_path
        self._ensure_file()

    def _ensure_file(self):
        if not os.path.exists(self.file_path):
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Signals"
            headers = ["Дата", "Матч", "Счет", "Время", "Общ голы", "Общ угловые",
                       "Общ удары", "Общ створ", "Общ ЖК",
                       "Хозяева голы", "Хозяева угловые", "Хозяева удары", "Хозяева створ", "Хозяева ЖК",
                       "Гости голы", "Гости угловые", "Гости удары", "Гости створ", "Гости ЖК",
                       "П1", "Ничья", "П2", "Тотал 2.5", "ID фильтра"]
            ws.append(headers)
            wb.save(self.file_path)

    def append_match(self, match: dict, stats: dict, odds: dict, filter_id: int):
        wb = load_workbook(self.file_path)
        ws = wb.active
        home_goals = stats.get('goals', {}).get('home', 0)
        away_goals = stats.get('goals', {}).get('away', 0)
        home_corners = stats.get('corners', {}).get('home', 0)
        away_corners = stats.get('corners', {}).get('away', 0)
        home_shots = stats.get('shots', {}).get('home', 0)
        away_shots = stats.get('shots', {}).get('away', 0)
        home_sot = stats.get('shots_on_target', {}).get('home', 0)
        away_sot = stats.get('shots_on_target', {}).get('away', 0)
        home_yellow = stats.get('yellow_cards', {}).get('home', 0)
        away_yellow = stats.get('yellow_cards', {}).get('away', 0)

        # Исправлено: используем homeTeam/awayTeam
        home_name = match.get('homeTeam', {}).get('name', '')
        away_name = match.get('awayTeam', {}).get('name', '')

        row = [
            datetime.now().isoformat(),
            f"{home_name} - {away_name}",
            f"{home_goals}-{away_goals}",
            match.get('elapsed', match.get('minute', 0)),
            home_goals + away_goals,
            home_corners + away_corners,
            home_shots + away_shots,
            home_sot + away_sot,
            home_yellow + away_yellow,
            home_goals, home_corners, home_shots, home_sot, home_yellow,
            away_goals, away_corners, away_shots, away_sot, away_yellow,
            odds.get('p1', 0), odds.get('draw', 0), odds.get('p2', 0), odds.get('total_over_2_5', 0),
            filter_id
        ]
        ws.append(row)
        wb.save(self.file_path)